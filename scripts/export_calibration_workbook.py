# -*- coding: utf-8 -*-
"""Build the calibration scoring workbook (.xlsx) for the human reviewers.

One workbook per reviewer slot, each with four sheets:
  1. 안내        -- what calibration is, how to score, the rules
  2. 채점 기준표 -- the 5-axis rubric with 0-33 / 34-66 / 67-100 anchors
  3. 채점 예시   -- worked examples (real hand-labelled utterances)
  4. 채점 시트   -- this reviewer's 90 blind items with empty score cells
                    (0-100 integer validation, frozen header)

Blind: the scoring sheet shows only the current utterance. Source, model
scores, the other reviewer's scores, and selection metadata are not included.

Requires openpyxl (dev-only tooling dependency, not a runtime import):
  .venv/Scripts/python.exe -m pip install openpyxl

Run from the repository root:
  python scripts/export_calibration_workbook.py
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ModuleNotFoundError as exc:  # pragma: no cover - environment guard
    raise SystemExit(
        "openpyxl is required: .venv/Scripts/python.exe -m pip install openpyxl"
    ) from exc

DEFAULT_MANIFEST = Path("data/fixtures/ml_transition_calibration_review.manifest.json")
DEFAULT_CANDIDATES = Path("data/fixtures/ml_transition_calibration_candidates_90.jsonl")
DEFAULT_OUTDIR = Path("data/fixtures")

AXES = ("Formality", "Energy", "Intimacy", "Humor", "Curiosity")
AXES_KO = {
    "Formality": "격식 (Formality)",
    "Energy": "에너지 (Energy)",
    "Intimacy": "친밀도 (Intimacy)",
    "Humor": "유머 (Humor)",
    "Curiosity": "호기심 (Curiosity)",
}

HEADER_FILL = PatternFill("solid", fgColor="2F5496")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14)
SUB_FONT = Font(bold=True, size=11)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


RUBRIC = {
    "Formality": {
        "정의": "말투가 얼마나 정중하고 조심스러운가. 학술/비즈니스 문어체가 아니라 '대화에서 공손하게 말하는 정도'.",
        "low": "반말/줄임말/친구 사이 인사. \"what's up\", \"gonna\", \"yeah no\".",
        "mid": "평범한 존중이 있는 대화체. 특별히 공손하지도 무례하지도 않음.",
        "high": "please / could you / would you / may I, 조심스럽고 예의 바른 구어체.",
        "포인트": "문장이 어려운 단어를 쓰는지가 아니라, 상대를 대하는 태도가 공손한지로 본다.",
    },
    "Energy": {
        "정의": "각성도와 감정의 세기. 차분함 ↔ 흥분/강조/감정 고조.",
        "low": "차분, 짧고 밋밋함, 감정 표시 거의 없음.",
        "mid": "보통 대화 톤. 약간의 강조나 감탄은 있으나 들뜨지 않음.",
        "high": "느낌표, 대문자 강조, \"so/really/very\", 감탄사, 빠른 리듬, 감정이 실림.",
        "포인트": "내용이 중요한지가 아니라 말하는 사람의 텐션으로 본다.",
    },
    "Intimacy": {
        "정의": "친밀함과 개인적 거리. 사무적/과제 중심 ↔ 친근/개인적/다정.",
        "low": "거리를 둔 사무적 발화, 과제나 정보만 다룸.",
        "mid": "약간의 개인적 표현(1인칭, 가벼운 근황)이 섞인 대화.",
        "high": "상대를 직접 부르거나, 개인적인 이야기, 공감/응원, 따뜻함.",
        "포인트": "말이 긴지가 아니라 관계의 거리감으로 본다.",
    },
    "Humor": {
        "정의": "유머와 장난기. 진지함 ↔ 농담/과장/밈/놀이적 표현.",
        "low": "문자 그대로, 진지함. 농담 의도 없음.",
        "mid": "가벼운 장난기나 미소 정도. 본격적인 농담은 아님.",
        "high": "명확한 농담, 과장, 자조, 말장난, 웃음(\"haha\", \"lol\") 동반.",
        "포인트": "웃긴 상황인지가 아니라 발화가 유머를 의도하는지로 본다.",
    },
    "Curiosity": {
        "정의": "궁금증과 질문성. 단정적 진술 ↔ 질문/궁금해함/설명 요청.",
        "low": "닫힌 단정 진술. 물음표 없음, 정보를 구하지 않음.",
        "mid": "가벼운 확인 질문(\"right?\")이나 약한 궁금증.",
        "high": "명확한 질문, \"how/why/what\", 설명·방법을 적극적으로 요청.",
        "포인트": "물음표 유무만 보지 말고, 무언가를 알고 싶어하는 의도를 본다.",
    },
}

SCALE_ROWS = [
    ("0-20", "신호 거의 없음 (해당 축 특성이 사실상 안 보임)"),
    ("21-40", "약한 신호 (조금 있으나 뚜렷하지 않음)"),
    ("41-60", "보통 신호 (분명히 있으나 지배적이지 않음)"),
    ("61-80", "강한 신호 (발화의 눈에 띄는 특징)"),
    ("81-100", "지배적 신호 (이 축이 발화를 규정함)"),
]

# Real hand-labelled seed utterances (data/dataset seed rows). Scores are the
# reference labels, the "이유" column is the short rationale for calibration.
EXAMPLES = [
    {
        "utterance": "Hey! How's your day treating you so far?",
        "Formality": 25, "Energy": 70, "Intimacy": 55, "Humor": 10, "Curiosity": 75,
        "이유": "친구에게 밝게 건네는 인사 겸 질문. 반말톤(격식 낮음), 느낌표+밝음(에너지 높음), 안부를 물음(호기심 높음).",
    },
    {
        "utterance": "Good afternoon, how have you been?",
        "Formality": 60, "Energy": 30, "Intimacy": 25, "Humor": 0, "Curiosity": 70,
        "이유": "같은 안부 질문이지만 정중한 인사체(격식 높음), 차분함(에너지 낮음), 거리감 있음(친밀도 낮음).",
    },
    {
        "utterance": "Could you please explain that one more time?",
        "Formality": 70, "Energy": 30, "Intimacy": 30, "Humor": 0, "Curiosity": 80,
        "이유": "please+could you로 공손함(격식 높음), 설명을 적극 요청(호기심 높음), 톤은 차분.",
    },
    {
        "utterance": "Wait, really? I've been saying it wrong for years! That's hilarious.",
        "Formality": 20, "Energy": 85, "Intimacy": 70, "Humor": 70, "Curiosity": 20,
        "이유": "놀람+느낌표 연발(에너지 높음), 자조적 농담 'hilarious'(유머 높음), 개인 경험 공유(친밀도 높음), 질문 아님(호기심 낮음).",
    },
    {
        "utterance": "In my opinion, learning languages is all about consistency.",
        "Formality": 50, "Energy": 50, "Intimacy": 35, "Humor": 0, "Curiosity": 10,
        "이유": "차분한 의견 진술. 단정문이라 호기심 낮음, 유머 없음, 톤/격식 모두 중간.",
    },
    {
        "utterance": "Um, is it 'went' or 'gone' here? I always forget.",
        "Formality": 30, "Energy": 20, "Intimacy": 20, "Humor": 5, "Curiosity": 60,
        "이유": "머뭇거림 'um'(에너지 낮음), 문법 확인 질문(호기심 중상), 캐주얼하지만 특별히 친근하진 않음.",
    },
    {
        "utterance": "I wanted to say... um, never mind, it's not that important anyway.",
        "Formality": 20, "Energy": 20, "Intimacy": 45, "Humor": 0, "Curiosity": 0,
        "이유": "말을 접는 자기수정. 낮은 텐션, 질문 아님(호기심 0), 개인적 머뭇거림이라 친밀도는 중간.",
    },
    {
        "utterance": "What's up, my friend? Anything exciting happening?",
        "Formality": 15, "Energy": 80, "Intimacy": 85, "Humor": 30, "Curiosity": 80,
        "이유": "매우 캐주얼('what's up')+호칭 'my friend'(친밀도 높음), 들뜬 톤(에너지 높음), 근황을 캐물음(호기심 높음).",
    },
]


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig") as handle:
        return [json.loads(line) for line in handle if line.strip()]


def _style_header(ws, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BOX


def sheet_intro(wb: Workbook, slot: str, item_count: int, seed: int) -> None:
    ws = wb.create_sheet("안내")
    ws.column_dimensions["A"].width = 100
    lines = [
        ("Pally 5축 Calibration 채점 안내", TITLE_FONT),
        ("", None),
        (f"검수자 슬롯: {slot}   /   문항 수: {item_count}   /   순서 seed: {seed}", SUB_FONT),
        ("", None),
        ("[ 목적 ]", SUB_FONT),
        ("두 검수자가 같은 기준으로 점수를 주는지 확인하는 단계입니다. 본격적인 대규모 라벨링 전에", None),
        ("기준 차이(한 사람이 계속 높게/낮게 주는지, 축 정의를 다르게 이해하는지)를 찾습니다.", None),
        ("이 데이터는 dev 용도이며 학습이나 최종 성능 판정에 쓰지 않습니다.", None),
        ("", None),
        ("[ 채점 방법 ]", SUB_FONT),
        ("1. '채점 기준표' 시트와 '채점 예시' 시트를 먼저 읽습니다.", None),
        ("2. '채점 시트'의 각 발화에 대해 5개 축을 각각 0~100 정수로 채웁니다.", None),
        ("3. 상단 reviewer_id 칸에 본인 식별자(예: chanhee)를 입력합니다.", None),
        ("4. 애매하거나 판단 근거가 특이한 경우 '메모' 칸에 한 줄 남깁니다.", None),
        ("5. 다른 검수자와 상의하지 말고 독립적으로 채점합니다.", None),
        ("", None),
        ("[ 규칙 ]", SUB_FONT),
        ("- 모든 점수는 0~100 범위. 소수점 없이 정수.", None),
        ("- 문어체가 아니라 '말로 한 문장'으로 보고 판단합니다.", None),
        ("- 학습자 실수, 짧은 조각, 자기수정, 말 멈춤도 대화면 정상입니다.", None),
        ("- 두 축이 충돌하면 둘 다 반영합니다. 예: 공손한 질문 = 격식 높음 + 호기심 높음.", None),
        ("- 현재 발화만 보고 판단합니다. 앞뒤 맥락/음성/화자 정보는 제공되지 않습니다.", None),
        ("- 발화 텍스트는 그대로 두고, 해석은 메모 칸에만 적습니다.", None),
        ("", None),
        ("[ 척도 감각 ]", SUB_FONT),
    ]
    r = 1
    for text, font in lines:
        cell = ws.cell(row=r, column=1, value=text)
        if font:
            cell.font = font
        cell.alignment = WRAP
        r += 1
    for band, meaning in SCALE_ROWS:
        ws.cell(row=r, column=1, value=f"  {band}  :  {meaning}").alignment = WRAP
        r += 1


def sheet_rubric(wb: Workbook) -> None:
    ws = wb.create_sheet("채점 기준표")
    headers = ["축", "정의", "낮음 (0-33)", "보통 (34-66)", "높음 (67-100)", "판단 포인트"]
    widths = [18, 44, 30, 30, 34, 40]
    for col, (head, width) in enumerate(zip(headers, widths), start=1):
        ws.cell(row=1, column=col, value=head)
        ws.column_dimensions[get_column_letter(col)].width = width
    _style_header(ws, 1, len(headers))
    for i, axis in enumerate(AXES, start=2):
        entry = RUBRIC[axis]
        values = [AXES_KO[axis], entry["정의"], entry["low"], entry["mid"], entry["high"], entry["포인트"]]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=i, column=col, value=value)
            cell.alignment = WRAP
            cell.border = BOX
        ws.row_dimensions[i].height = 118
    note_row = len(AXES) + 3
    ws.cell(row=note_row, column=1,
            value="0-33 / 34-66 / 67-100 은 큰 구간 감각용입니다. 실제 점수는 그 구간 안에서 신호의 세기에 따라 자유롭게 정하세요 (위 '척도 감각' 참고).").alignment = WRAP
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=6)


def sheet_examples(wb: Workbook) -> None:
    ws = wb.create_sheet("채점 예시")
    headers = ["예시 발화", *[AXES_KO[a] for a in AXES], "채점 이유"]
    widths = [46, 13, 13, 13, 13, 13, 60]
    for col, (head, width) in enumerate(zip(headers, widths), start=1):
        ws.cell(row=1, column=col, value=head)
        ws.column_dimensions[get_column_letter(col)].width = width
    _style_header(ws, 1, len(headers))
    for i, ex in enumerate(EXAMPLES, start=2):
        ws.cell(row=i, column=1, value=ex["utterance"]).alignment = WRAP
        for j, axis in enumerate(AXES, start=2):
            cell = ws.cell(row=i, column=j, value=ex[axis])
            cell.alignment = CENTER
        ws.cell(row=i, column=7, value=ex["이유"]).alignment = WRAP
        for col in range(1, 8):
            ws.cell(row=i, column=col).border = BOX
        ws.row_dimensions[i].height = 74


def sheet_scoring(wb: Workbook, slot: str, items: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("채점 시트")
    ws.cell(row=1, column=1, value="reviewer_id →").font = SUB_FONT
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="right")
    ws.cell(row=2, column=1, value="reviewer_slot →").font = SUB_FONT
    ws.cell(row=2, column=1).alignment = Alignment(horizontal="right")
    ws.cell(row=2, column=2, value=slot)

    header_row = 4
    headers = ["item_id", "발화 (utterance)", *[AXES_KO[a] for a in AXES], "메모 (선택)"]
    widths = [16, 62, 12, 12, 12, 12, 12, 40]
    for col, (head, width) in enumerate(zip(headers, widths), start=1):
        ws.cell(row=header_row, column=col, value=head)
        ws.column_dimensions[get_column_letter(col)].width = width
    _style_header(ws, header_row, len(headers))

    validation = DataValidation(type="whole", operator="between", formula1=0, formula2=100, allow_blank=True)
    validation.error = "0에서 100 사이 정수만 입력하세요."
    validation.errorTitle = "범위 초과"
    validation.prompt = "0~100 정수"
    ws.add_data_validation(validation)

    for offset, item in enumerate(items):
        r = header_row + 1 + offset
        ws.cell(row=r, column=1, value=item["item_id"]).alignment = CENTER
        ws.cell(row=r, column=2, value=item["utterance"]).alignment = WRAP
        for col in range(3, 8):
            cell = ws.cell(row=r, column=col)
            cell.alignment = CENTER
            cell.border = BOX
            validation.add(cell)
        ws.cell(row=r, column=8).alignment = WRAP
        for col in (1, 2, 8):
            ws.cell(row=r, column=col).border = BOX
        ws.row_dimensions[r].height = 30

    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)


def build_workbook(slot: str, items: list[dict[str, Any]], seed: int, out_path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    sheet_intro(wb, slot, len(items), seed)
    sheet_rubric(wb)
    sheet_examples(wb)
    sheet_scoring(wb, slot, items)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--candidates", type=Path, default=DEFAULT_CANDIDATES)
    parser.add_argument("--outdir", type=Path, default=DEFAULT_OUTDIR)
    return parser


def main() -> None:
    args = build_parser().parse_args()
    manifest = json.loads(args.manifest.read_text(encoding="utf-8"))
    candidates = load_jsonl(args.candidates)
    by_index = {index: row for index, row in enumerate(candidates)}
    item_by_id = {item["item_id"]: item for item in manifest["items"]}

    written = []
    for slot in manifest["slots"]:
        ordered_ids = manifest["slot_order"][slot]
        items = []
        for item_id in ordered_ids:
            manifest_item = item_by_id[item_id]
            items.append(
                {
                    "item_id": item_id,
                    "utterance": manifest_item["utterance"],
                }
            )
        out_path = args.outdir / f"calibration_scoring_slot{slot}.xlsx"
        build_workbook(slot, items, manifest["seed"], out_path)
        written.append(out_path)

    for path in written:
        print(f"wrote {path}")
    print(f"items per slot: {manifest['item_count']}, slots: {','.join(manifest['slots'])}")


if __name__ == "__main__":
    main()
